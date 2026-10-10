from datetime import datetime
from pathlib import Path
import traceback as traceback_module

from openpyxl import Workbook, load_workbook

from src.configs.path import DOWNLOAD_DIR, RESULT_DIR

RESULT_HEADERS = [
    "post_order",
    "via_name",
    "page_name",
    "caption",
    "image",
    "comment",
    "post_status",
    "published_at",
    "post_url",
    "error",
    "_task_id",
]

DEVLOG_HEADERS = [
    "timestamp",
    "run_id",
    "post_order",
    "via_id",
    "profile_id",
    "page_id",
    "action",
    "status",
    "attempt",
    "error_type",
    "error_message",
    "traceback",
    "post_url",
]


def _now():
    return datetime.now().astimezone().replace(tzinfo=None)


class PageAutoReport:
    def __init__(self):
        self.result_path = Path(DOWNLOAD_DIR) / "PageAuto_Result.xlsx"
        self.devlog_path = Path(RESULT_DIR) / "PageAuto_DevLog.xlsx"

        self._ensure_result_file()
        self._ensure_devlog_file()
        self.results = self._read_results()

    def _ensure_result_file(self):
        if self.result_path.exists():
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"
        sheet.append(RESULT_HEADERS)
        sheet.column_dimensions["K"].hidden = True
        workbook.save(self.result_path)
        workbook.close()

    def _ensure_devlog_file(self):
        if self.devlog_path.exists():
            workbook = load_workbook(self.devlog_path)

            if "DevLog" not in workbook.sheetnames:
                workbook.create_sheet("DevLog").append(DEVLOG_HEADERS)

            if "RUN_STATE" not in workbook.sheetnames:
                workbook.create_sheet("RUN_STATE").append(["key", "value"])

            workbook.save(self.devlog_path)
            workbook.close()
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "DevLog"
        sheet.append(DEVLOG_HEADERS)

        state = workbook.create_sheet("RUN_STATE")
        state.append(["key", "value"])

        workbook.save(self.devlog_path)
        workbook.close()

    def _read_results(self) -> dict:
        workbook = load_workbook(
            self.result_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook["Results"]
            headers = [cell.value for cell in sheet[1]]
            positions = {name: index for index, name in enumerate(headers)}
            results = {}

            for values in sheet.iter_rows(min_row=2, values_only=True):
                task_id = (
                    values[positions["_task_id"]] if "_task_id" in positions else None
                )

                if not task_id:
                    continue

                results[str(task_id)] = {
                    "post_status": (
                        values[positions["post_status"]]
                        if "post_status" in positions
                        else ""
                    ),
                    "published_at": (
                        values[positions["published_at"]]
                        if "published_at" in positions
                        else None
                    ),
                    "post_url": (
                        values[positions["post_url"]] if "post_url" in positions else ""
                    ),
                    "error": (
                        values[positions["error"]] if "error" in positions else ""
                    ),
                }

            return results
        finally:
            workbook.close()

    def get_result(self, task_id: str) -> dict:
        return self.results.get(str(task_id), {}).copy()

    def upsert_result(
        self,
        row: dict,
        status: str,
        published_at=None,
        post_url="",
        error="",
    ) -> None:
        task_id = str(row.get("task_id") or "")
        if not task_id:
            raise ValueError("Cannot save a result without task_id.")

        workbook = load_workbook(self.result_path)

        try:
            sheet = workbook["Results"]
            headers = [cell.value for cell in sheet[1]]

            if "_task_id" not in headers:
                raise ValueError("PageAuto_Result.xlsx is missing the _task_id column.")

            task_column = headers.index("_task_id") + 1
            target_row = None

            for row_index in range(2, sheet.max_row + 1):
                value = sheet.cell(row_index, task_column).value
                if str(value or "") == task_id:
                    target_row = row_index
                    break

            if target_row is None:
                target_row = sheet.max_row + 1

            comment = str(row.get("comment") or "").strip()
            if not comment:
                comment = "No comment (skipped)"

            values = {
                "post_order": row.get("post_order"),
                "via_name": row.get("via_name", ""),
                "page_name": row.get("page", row.get("page_name", "")),
                "caption": row.get("caption", ""),
                "image": row.get("image", ""),
                "comment": comment,
                "post_status": status,
                "published_at": published_at,
                "post_url": post_url or "",
                "error": error or "",
                "_task_id": task_id,
            }

            for column, header in enumerate(headers, start=1):
                if header in values:
                    sheet.cell(target_row, column, values[header])

            sheet.column_dimensions["K"].hidden = True
            workbook.save(self.result_path)

            self.results[task_id] = {
                "post_status": status,
                "published_at": published_at,
                "post_url": post_url or "",
                "error": error or "",
            }
        finally:
            workbook.close()

    def log_event(
        self,
        row: dict,
        run_id: str,
        action: str,
        status: str,
        attempt: int = 0,
        error=None,
        post_url="",
    ) -> None:
        error_type = type(error).__name__ if error else ""
        error_message = str(error) if error else ""
        trace = ""

        if error:
            trace = "".join(
                traceback_module.format_exception(
                    type(error),
                    error,
                    error.__traceback__,
                )
            )

        values = [
            _now(),
            run_id,
            row.get("post_order"),
            row.get("via_id", ""),
            row.get("profile_id", ""),
            row.get("page_id", ""),
            action,
            status,
            attempt,
            error_type,
            error_message,
            trace,
            post_url or "",
        ]

        workbook = load_workbook(self.devlog_path)

        try:
            workbook["DevLog"].append(values)
            workbook.save(self.devlog_path)
        finally:
            workbook.close()

    def next_attempt(self, row: dict) -> int:
        action = f"publish:{row.get('post_id', '')}"

        workbook = load_workbook(
            self.devlog_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook["DevLog"]
            headers = [cell.value for cell in sheet[1]]
            action_index = headers.index("action")
            attempt_index = headers.index("attempt")
            highest = 0

            for values in sheet.iter_rows(min_row=2, values_only=True):
                if values[action_index] != action:
                    continue

                try:
                    highest = max(
                        highest,
                        int(values[attempt_index] or 0),
                    )
                except (TypeError, ValueError):
                    pass

            return highest + 1
        finally:
            workbook.close()

    def get_state(self) -> dict:
        workbook = load_workbook(
            self.devlog_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook["RUN_STATE"]
            return {
                str(key): value
                for key, value in sheet.iter_rows(
                    min_row=2,
                    values_only=True,
                )
                if key
            }
        finally:
            workbook.close()

    def set_state(self, state: dict) -> None:
        workbook = load_workbook(self.devlog_path)

        try:
            sheet = workbook["RUN_STATE"]
            current_rows = {}

            for row_index in range(2, sheet.max_row + 1):
                key = sheet.cell(row_index, 1).value
                if key:
                    current_rows[str(key)] = row_index

            for key, value in state.items():
                row_index = current_rows.get(key)

                if row_index is None:
                    row_index = sheet.max_row + 1

                sheet.cell(row_index, 1, key)
                sheet.cell(row_index, 2, value)

            workbook.save(self.devlog_path)
        finally:
            workbook.close()
