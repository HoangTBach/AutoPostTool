from pathlib import Path

from openpyxl import load_workbook

from src.configs.path import DATA_FILE
from src.features.page_posting.helper import clean_value, get_number


# ----- File -----
def get_data_file_stamp() -> tuple[int, int]:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Cannot find: {DATA_FILE}")

    stat = DATA_FILE.stat()

    return stat.st_mtime_ns, stat.st_size


# ----- Load -----
def load_page_rows() -> list[dict]:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Cannot find: {DATA_FILE}")

    workbook = load_workbook(
        DATA_FILE,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    try:
        profiles = _sheet_to_dicts(workbook["ADSPOWER"])
        vias = _sheet_to_dicts(workbook["VIA"])
        pages = _sheet_to_dicts(workbook["PAGE"])
        posts = _sheet_to_dicts(workbook["PAGE_POST"])

        profile_map = _build_map(profiles, "profile_id")
        via_map = _build_map(vias, "via_id")
        page_map = _build_map(pages, "page_id")

        rows = []

        for post in posts:
            post_id = clean_value(post.get("post_id"))
            page_id = clean_value(post.get("page_id"))

            if not post_id or not page_id:
                continue

            page = page_map.get(page_id)

            if not page:
                continue

            via_id = clean_value(page.get("via_id"))
            via = via_map.get(via_id)

            if not via:
                continue

            profile_id = clean_value(via.get("profile_id"))
            profile = profile_map.get(profile_id, {})

            rows.append(
                {
                    "task_id": f"{post_id}|{profile_id}|{via_id}|{page_id}",
                    "post_id": post_id,
                    "post_order": get_number(post.get("post_order")),
                    "profile_id": profile_id,
                    "profile_uid": clean_value(profile.get("profile_uid")),
                    "profile_name": clean_value(profile.get("profile_name")),
                    "profile_status": clean_value(profile.get("status")),
                    "via_id": via_id,
                    "via_name": clean_value(via.get("via_name")),
                    "page_id": page_id,
                    "page": clean_value(page.get("page_name")) or page_id,
                    "page_link": clean_value(page.get("page_link")),
                    "page_status": clean_value(page.get("status")),
                    "caption": clean_value(post.get("caption")),
                    "image": clean_value(post.get("image")),
                    "comment": clean_value(post.get("comment")),
                    "status": "",
                }
            )

        return rows

    finally:
        workbook.close()


# ----- Map -----
def _build_map(rows: list[dict], key: str) -> dict[str, dict]:
    result = {}

    for row in rows:
        value = clean_value(row.get(key))

        if not value:
            continue

        if value in result:
            raise ValueError(f"Duplicate {key}: {value}")

        result[value] = row

    return result


# ----- Sheet -----
def _sheet_to_dicts(sheet) -> list[dict]:
    values = sheet.iter_rows(values_only=True)

    try:
        headers = next(values)
    except StopIteration:
        return []

    headers = [clean_value(header) for header in headers]
    rows = []

    for values_row in values:
        if not any(clean_value(value) for value in values_row):
            continue

        row = {}

        for index, header in enumerate(headers):
            if not header:
                continue

            row[header] = values_row[index] if index < len(values_row) else None

        rows.append(row)

    return rows
